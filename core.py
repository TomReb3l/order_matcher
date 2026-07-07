from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd
import pdfplumber
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TABLE_COLUMNS = [
    "Α/Α PDF",
    "Μητρώο",
    "Επώνυμο",
    "Όνομα",
    "Πατρώνυμο",
    "Βαθμός",
    "Οργανική",
    "Τρόπος Ταύτισης",
]


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""

    text = str(value).strip().upper()
    text = " ".join(text.split())
    text = text.replace("Ϊ", "Ι").replace("Ϋ", "Υ")
    text = "".join(ch for ch in unicodedata.normalize("NFD", text) if unicodedata.category(ch) != "Mn")
    text = re.sub(r"[^A-ZΑ-Ω0-9 ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def safe_registry(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    match = re.search(r"(\d{4,})", text)
    return match.group(1) if match else ""


class PromotionPdfParser:
    SURNAME_X_MAX = 245
    NAME_X_MAX = 338
    PATRONYMIC_X_MAX = 468

    NOISE_TOKENS = {
        "ΥΠ.",
        "ΓΡΑΦ.",
        "ΥΠ",
        "ΓΡΑΦ",
        "ΥΠ.ΓΡΑΦ.",
        "ΥΠ. ΓΡΑΦ.",
    }

    def parse(self, pdf_path: str | Path) -> pd.DataFrame:
        rows: list[dict] = []

        with pdfplumber.open(str(pdf_path)) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                words = page.extract_words(
                    x_tolerance=1,
                    y_tolerance=3,
                    keep_blank_chars=False,
                    use_text_flow=False,
                )
                grouped_by_line: dict[float, list[dict]] = defaultdict(list)
                for word in words:
                    grouped_by_line[round(word["top"], 1)].append(word)

                for top in sorted(grouped_by_line):
                    line_words = sorted(grouped_by_line[top], key=lambda item: item["x0"])
                    texts = [w["text"] for w in line_words]
                    if len(texts) < 2:
                        continue

                    # FEK/PDF tables usually render the row number as "1.", "142." etc.
                    # The previous parser accepted only bare digits ("142"), so it skipped
                    # valid rows and could miss registry matches entirely.
                    row_no_match = re.fullmatch(r"(\d{1,4})\.?", texts[0].strip())
                    registry = texts[1].strip()

                    if not (row_no_match and re.fullmatch(r"\d{6}", registry)):
                        continue

                    record = {
                        "pdf_page": page_index,
                        "pdf_row_no": int(row_no_match.group(1)),
                        "registry": registry,
                        "surname": [],
                        "name": [],
                        "patronymic": [],
                        "extra": [],
                    }

                    for word in line_words[2:]:
                        x = word["x0"]
                        token = word["text"].strip()
                        if not token:
                            continue

                        if x < self.SURNAME_X_MAX:
                            record["surname"].append(token)
                        elif x < self.NAME_X_MAX:
                            record["name"].append(token)
                        elif x < self.PATRONYMIC_X_MAX:
                            record["patronymic"].append(token)
                        else:
                            record["extra"].append(token)

                    surname = " ".join(record["surname"]).strip()
                    name = " ".join(record["name"]).strip()
                    patronymic_tokens = [t for t in record["patronymic"] if normalize_text(t) not in self.NOISE_TOKENS]
                    extra_tokens = [t for t in record["extra"] if normalize_text(t) not in self.NOISE_TOKENS]
                    patronymic = " ".join(patronymic_tokens).strip()
                    extra = " ".join(extra_tokens).strip()

                    if not surname or not name:
                        continue

                    rows.append(
                        {
                            "pdf_page": record["pdf_page"],
                            "pdf_row_no": record["pdf_row_no"],
                            "registry": record["registry"],
                            "surname": surname,
                            "name": name,
                            "patronymic": patronymic,
                            "extra": extra,
                        }
                    )

        df = pd.DataFrame(rows)
        if df.empty:
            raise ValueError("Δεν βρέθηκαν εγγραφές στο PDF.")

        df = df.drop_duplicates(subset=["registry", "surname", "name", "patronymic"]).reset_index(drop=True)
        for col in ["registry", "surname", "name", "patronymic"]:
            df[f"norm_{col}"] = df[col].map(normalize_text)
        return df


class ServiceExcelLoader:
    REQUIRED_COLUMNS = {"ΜΗΤΡΩΟ", "ΕΠΩΝΥΜΟ", "ΟΝΟΜΑ"}

    def load(self, excel_path: str | Path) -> pd.DataFrame:
        xl = pd.ExcelFile(excel_path)
        best_error = None
        for sheet_name in xl.sheet_names:
            try:
                return self._load_sheet(excel_path, sheet_name)
            except Exception as exc:  # noqa: BLE001
                best_error = exc
        if best_error:
            raise best_error
        raise ValueError("Δεν ήταν δυνατή η φόρτωση του Excel.")

    def _load_sheet(self, excel_path: str | Path, sheet_name: str) -> pd.DataFrame:
        raw = pd.read_excel(excel_path, sheet_name=sheet_name, header=None)
        header_row_index = self._find_header_row(raw)
        if header_row_index is None:
            raise ValueError(f"Δεν βρέθηκε γραμμή επικεφαλίδων στο φύλλο: {sheet_name}")

        header_values = [normalize_text(v) for v in raw.iloc[header_row_index].tolist()]
        df = raw.iloc[header_row_index + 1 :].copy()
        df.columns = header_values
        df = df.dropna(how="all").reset_index(drop=True)

        keep_cols = [c for c in df.columns if c]
        df = df[keep_cols]

        registry_col = self._pick_column(df.columns, ["ΜΗΤΡΩΟ", "ΑΡΙΘΜΟΣ ΜΗΤΡΩΟΥ", "Α Μ", "ΑΜ"])
        surname_col = self._pick_column(df.columns, ["ΕΠΩΝΥΜΟ"])
        name_col = self._pick_column(df.columns, ["ΟΝΟΜΑ"])
        patronymic_col = self._pick_column(df.columns, ["ΠΑΤΡΩΝΥΜΟ"])
        rank_col = self._pick_column(df.columns, ["ΒΑΘΜΟΣ"]) or ""
        unit_col = self._pick_column(df.columns, ["ΟΡΓΑΝΙΚΗ", "ΥΠΗΡΕΣΙΑ", "ΤΜΗΜΑ", "ΔΙΕΥΘΥΝΣΗ"]) or ""

        if not registry_col or not surname_col or not name_col:
            raise ValueError("Δεν εντοπίστηκαν οι βασικές στήλες Μητρώο / Επώνυμο / Όνομα στο Excel.")

        out = pd.DataFrame(
            {
                "registry": df[registry_col].map(safe_registry),
                "rank": df[rank_col].astype(str).replace("nan", "").str.strip() if rank_col else "",
                "surname": df[surname_col].astype(str).replace("nan", "").str.strip(),
                "name": df[name_col].astype(str).replace("nan", "").str.strip(),
                "patronymic": df[patronymic_col].astype(str).replace("nan", "").str.strip() if patronymic_col else "",
                "service_unit": df[unit_col].astype(str).replace("nan", "").str.strip() if unit_col else "",
                "source_sheet": sheet_name,
            }
        )
        out = out[out["registry"].astype(str).str.strip() != ""].copy()
        out = out.reset_index(drop=True)

        for col in ["registry", "surname", "name", "patronymic"]:
            out[f"norm_{col}"] = out[col].map(normalize_text)
        return out

    def _find_header_row(self, raw: pd.DataFrame) -> Optional[int]:
        for idx, row in raw.iterrows():
            normalized = {normalize_text(v) for v in row.tolist()}
            if self.REQUIRED_COLUMNS.issubset(normalized):
                return int(idx)
        return None

    @staticmethod
    def _pick_column(columns: list[str] | pd.Index, candidates: list[str]) -> Optional[str]:
        normalized_columns = [normalize_text(c) for c in columns]
        for candidate in candidates:
            target = normalize_text(candidate)
            for col, norm_col in zip(columns, normalized_columns):
                if norm_col == target:
                    return str(col)
        for candidate in candidates:
            target = normalize_text(candidate)
            for col, norm_col in zip(columns, normalized_columns):
                if target and target in norm_col:
                    return str(col)
        return None


class MatcherEngine:
    def match(self, promotions_df: pd.DataFrame, service_df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
        if promotions_df.empty:
            raise ValueError("Το PDF δεν έδωσε εγγραφές για σύγκριση.")
        if service_df.empty:
            raise ValueError("Το Excel δεν έδωσε εγγραφές για σύγκριση.")

        exact_registry = promotions_df.merge(
            service_df,
            left_on="norm_registry",
            right_on="norm_registry",
            how="inner",
            suffixes=("_pdf", "_excel"),
        ).copy()
        exact_registry["match_method"] = "Μητρώο"
        exact_registry["name_key"] = (
            exact_registry["norm_surname_pdf"]
            + "|"
            + exact_registry["norm_name_pdf"]
            + "|"
            + exact_registry["norm_patronymic_pdf"]
        )

        matched_registry_keys = set(exact_registry["norm_registry"].tolist())
        pdf_left = promotions_df[~promotions_df["norm_registry"].isin(matched_registry_keys)].copy()
        excel_left = service_df[~service_df["norm_registry"].isin(matched_registry_keys)].copy()

        pdf_left["name_key"] = (
            pdf_left["norm_surname"] + "|" + pdf_left["norm_name"] + "|" + pdf_left["norm_patronymic"]
        )
        excel_left["name_key"] = (
            excel_left["norm_surname"] + "|" + excel_left["norm_name"] + "|" + excel_left["norm_patronymic"]
        )

        exact_names = pdf_left.merge(
            excel_left,
            on="name_key",
            how="inner",
            suffixes=("_pdf", "_excel"),
        ).copy()
        if not exact_names.empty:
            exact_names["match_method"] = "Ονοματεπώνυμο + Πατρώνυμο"
            exact_names["norm_registry"] = exact_names["norm_registry_pdf"]

        common = pd.concat([exact_registry, exact_names], ignore_index=True, sort=False)
        common = common.drop_duplicates(subset=["registry_pdf", "registry_excel", "name_key"], keep="first")

        result = pd.DataFrame(
            {
                "Α/Α PDF": common.get("pdf_row_no", ""),
                "Μητρώο": common.get("registry_pdf", common.get("registry_excel", "")),
                "Επώνυμο": common.get("surname_pdf", common.get("surname_excel", "")),
                "Όνομα": common.get("name_pdf", common.get("name_excel", "")),
                "Πατρώνυμο": common.get("patronymic_pdf", common.get("patronymic_excel", "")),
                "Βαθμός": common.get("rank", ""),
                "Οργανική": common.get("service_unit", ""),
                "Τρόπος Ταύτισης": common.get("match_method", ""),
                "PDF Σελίδα": common.get("pdf_page", ""),
                "Excel Φύλλο": common.get("source_sheet", ""),
            }
        )
        result = result.sort_values(by=["Α/Α PDF", "Επώνυμο", "Όνομα"], na_position="last").reset_index(drop=True)

        found_registry = set(result["Μητρώο"].astype(str).tolist())
        only_promotions = promotions_df[~promotions_df["registry"].isin(found_registry)].copy()
        only_service = service_df[~service_df["registry"].isin(found_registry)].copy()

        summary = {
            "promotions_total": int(len(promotions_df)),
            "service_total": int(len(service_df)),
            "common_total": int(len(result)),
            "registry_matches": int((result["Τρόπος Ταύτισης"] == "Μητρώο").sum()) if not result.empty else 0,
            "name_matches": int((result["Τρόπος Ταύτισης"] == "Ονοματεπώνυμο + Πατρώνυμο").sum()) if not result.empty else 0,
            "only_promotions_total": int(len(only_promotions)),
            "only_service_total": int(len(only_service)),
        }

        return result, {
            "common": result,
            "only_promotions": only_promotions,
            "only_service": only_service,
            "summary": summary,
        }


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


class WordExporter:
    def export(self, output_path: str | Path, common_df: pd.DataFrame, summary: dict, source_pdf: str, source_excel: str) -> None:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.3)
        section.left_margin = Cm(1.05)
        section.right_margin = Cm(1.05)
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

        styles = doc.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(8.5)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(3)
        run = title.add_run("Κοινοί μεταξύ Διαταγής και Δύναμης Υπηρεσίας")
        run.bold = True
        run.font.size = Pt(15)

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(9)
        subtitle_run = subtitle.add_run(datetime.now().strftime("Ημερομηνία εξαγωγής: %d/%m/%Y %H:%M"))
        subtitle_run.italic = True
        subtitle_run.font.size = Pt(8.5)

        meta = doc.add_table(rows=4, cols=2)
        meta.style = "Table Grid"
        meta.autofit = False
        set_table_fixed_layout(meta)
        meta_widths = [Cm(4.1), Cm(22.0)]
        meta_rows = [
            ("PDF Διαταγής", str(source_pdf)),
            ("Excel Δύναμης", str(source_excel)),
            ("Σύνολο κοινών", str(summary.get("common_total", 0))),
            (
                "Τρόπος ταύτισης",
                f"Μητρώο: {summary.get('registry_matches', 0)} | Ονοματεπώνυμο + Πατρώνυμο: {summary.get('name_matches', 0)}",
            ),
        ]
        for i, (left, right) in enumerate(meta_rows):
            for j, value in enumerate((left, right)):
                cell = meta.cell(i, j)
                cell.width = meta_widths[j]
                cell.text = str(value)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=55, start=95, bottom=55, end=95)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if paragraph.runs:
                    paragraph.runs[0].font.size = Pt(8.5)
            if meta.cell(i, 0).paragraphs[0].runs:
                meta.cell(i, 0).paragraphs[0].runs[0].bold = True

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(5)

        table = doc.add_table(rows=1, cols=len(TABLE_COLUMNS))
        table.style = "Table Grid"
        table.autofit = False
        set_table_fixed_layout(table)
        widths_cm = [1.05, 1.55, 3.35, 3.05, 2.8, 1.55, 5.6, 4.1]
        header_row = table.rows[0]
        for idx, col_name in enumerate(TABLE_COLUMNS):
            cell = header_row.cells[idx]
            cell.width = Cm(widths_cm[idx])
            cell.text = col_name
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=65, start=85, bottom=65, end=85)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if paragraph.runs:
                paragraph.runs[0].bold = True
                paragraph.runs[0].font.size = Pt(8)
        set_repeat_table_header(header_row)

        if common_df.empty:
            row = table.add_row()
            row.cells[0].merge(row.cells[-1])
            row.cells[0].text = "Δεν βρέθηκαν κοινά άτομα."
            row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_margins(row.cells[0], top=70, start=100, bottom=70, end=100)
        else:
            for _, item in common_df.iterrows():
                row = table.add_row()
                values = [
                    item.get("Α/Α PDF", ""),
                    item.get("Μητρώο", ""),
                    item.get("Επώνυμο", ""),
                    item.get("Όνομα", ""),
                    item.get("Πατρώνυμο", ""),
                    item.get("Βαθμός", ""),
                    item.get("Οργανική", ""),
                    item.get("Τρόπος Ταύτισης", ""),
                ]
                for idx, value in enumerate(values):
                    cell = row.cells[idx]
                    cell.width = Cm(widths_cm[idx])
                    cell.text = "" if pd.isna(value) else str(value)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_margins(cell, top=55, start=85, bottom=55, end=85)
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 1, 5) else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        run.font.size = Pt(8)

        doc.save(str(output_path))
